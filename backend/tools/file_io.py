"""
File I/O Tool: Path-scoped read/write for agent file operations.
All operations are sandboxed to the workspace directory (sandbox_files/ and outputs/).
Directory traversal outside the workspace is blocked to prevent escape.
"""

import os
import re
from pathlib import Path

# Image extensions whose contents are never usable as LLM text input.
_IMAGE_EXTENSIONS = {
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff",
    ".ico", ".heic", ".svg", ".avif",
}

# Workspace root directory
WORKSPACE_DIR = (Path(__file__).parent.parent.parent / "workspace").resolve()

# Base directory for agent input files and sandbox files
BASE_DIR = WORKSPACE_DIR / "sandbox_files"
BASE_DIR.mkdir(parents=True, exist_ok=True)

# Deliverables and generated reports directory
OUTPUT_DIR = WORKSPACE_DIR / "outputs"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def _safe_resolve(filename: str, for_write: bool = False) -> Path:
    """
    Resolve filename within the workspace sandbox with traversal protection.
    Normalizes absolute paths, relative paths, and home directory '~'.
    Checks BASE_DIR first, and falls back to OUTPUT_DIR for read operations.
    Raises ValueError if the resolved path escapes the sandbox / workspace.
    """
    if not filename or not str(filename).strip():
        raise ValueError("Empty filename provided.")

    filename_str = str(filename).strip()

    if for_write:
        # Writes are strictly sandboxed to BASE_DIR
        target = (BASE_DIR / filename_str).resolve()
        try:
            target.relative_to(BASE_DIR)
        except ValueError:
            raise ValueError(
                f"Path traversal detected: '{filename}' resolves outside the sandbox."
            )
        return target

    # Normalize home-directory paths if provided
    if filename_str.startswith("~"):
        p = Path(filename_str).expanduser().resolve()
    else:
        p = Path(filename_str)

    if p.is_absolute():
        target = p.resolve()
        try:
            target.relative_to(WORKSPACE_DIR)
        except ValueError:
            raise ValueError(
                f"Path traversal detected: '{filename}' resolves outside the workspace."
            )
        return target

    # Handle relative paths: normalize leading workspace/ or ./
    norm = filename_str
    if norm.startswith("./") or norm.startswith(".\\"):
        norm = norm[2:]

    if norm.startswith("workspace/") or norm.startswith("workspace\\"):
        norm = norm[len("workspace/"):]

    # If the user explicitly provided outputs/ or sandbox_files/ prefix
    if norm.startswith("outputs/") or norm.startswith("outputs\\"):
        target = (WORKSPACE_DIR / norm).resolve()
    elif norm.startswith("sandbox_files/") or norm.startswith("sandbox_files\\"):
        target = (WORKSPACE_DIR / norm).resolve()
    else:
        # Default: check BASE_DIR first
        target = (BASE_DIR / norm).resolve()
        # For reads: if file does not exist in BASE_DIR, check OUTPUT_DIR fallback
        if not target.exists():
            candidate = (OUTPUT_DIR / norm).resolve()
            if candidate.exists():
                target = candidate

    # Ensure resolved path is strictly within WORKSPACE_DIR
    try:
        target.relative_to(WORKSPACE_DIR)
    except ValueError:
        raise ValueError(
            f"Path traversal detected: '{filename}' resolves outside the workspace."
        )

    return target


def read_file(filename: str) -> str:
    """
    Read a file from the sandboxed directory or deliverables outputs.
    Supports .txt, .md, .csv, .json, .py, .docx, .pptx, .xlsx, .pdf.

    Args:
        filename: Relative or workspace path (no directory traversal allowed).

    Returns:
        File contents as a string, or an error message.
    """
    try:
        path = _safe_resolve(filename, for_write=False)
    except ValueError as e:
        return f"Error: {e}"

    # Resolve to the real on-disk file when the caller's casing differs from
    # the actual filename (uploaded files keep their original mixed case).
    # Exact-name matches are always preferred; a unique case-variant match is
    # the only fallback, so we never silently pick between ambiguous files.
    from backend.tools.path_safety import resolve_existing_casefold
    real_path = resolve_existing_casefold(path)
    if real_path is not None:
        path = real_path

    if not path.exists():
        return f"Error: File not found: {filename}"

    if not path.is_file():
        return f"Error: Path is not a file: {filename}"

    # Image and other binary files must NEVER be decoded as raw text and fed to
    # an LLM: a ~40 KB screenshot decodes to ~30k garbage tokens, which blows
    # past the model context window ("Requested tokens (31520) exceed context
    # window of 2048"). Return a descriptive placeholder pointing at the image
    # analysis tools instead.
    suffix = path.suffix.lower()
    if suffix in _IMAGE_EXTENSIONS:
        return (
            f"[Image file: {Path(path).name}] This is a binary image ({path.stat().st_size} bytes), "
            "not text — it cannot be read or summarized as text content. "
            f"To analyze the image, ask e.g. \"Analyze the nameplate in {path.name}\", "
            f"\"Read the handwriting in {path.name}\", or \"Extract the P&ID from {path.name}\"."
        )

    try:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            try:
                import pypdf
                reader = pypdf.PdfReader(str(path))
                pages = [page.extract_text() or "" for page in reader.pages]
                text = "\n\n".join(pages).strip()
                if not text:
                    return f"[Empty or image-based PDF: {filename}]"
                return text
            except Exception as pe:
                return f"Error extracting PDF text: {pe}"

        elif suffix == ".docx":
            try:
                import docx
                doc = docx.Document(str(path))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                table_lines = []
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells)
                        if row_text.strip():
                            table_lines.append(row_text)
                full_text = "\n".join(paragraphs)
                if table_lines:
                    full_text += "\n\nTables:\n" + "\n".join(table_lines)
                return full_text if full_text.strip() else f"[Empty docx: {filename}]"
            except Exception as de:
                return f"Error reading docx file: {de}"

        elif suffix == ".pptx":
            try:
                import pptx
                prs = pptx.Presentation(str(path))
                slide_texts = []
                for idx, slide in enumerate(prs.slides, 1):
                    shapes_text = []
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for p in shape.text_frame.paragraphs:
                                if p.text.strip():
                                    shapes_text.append(p.text.strip())
                        elif shape.has_table:
                            for row in shape.table.rows:
                                row_str = " | ".join(cell.text.strip() for cell in row.cells)
                                if row_str.strip():
                                    shapes_text.append(row_str)
                    if shapes_text:
                        slide_texts.append(f"Slide {idx}:\n" + "\n".join(shapes_text))
                return "\n\n".join(slide_texts).strip() if slide_texts else f"[Empty presentation: {filename}]"
            except Exception as pe:
                return f"Error reading pptx file: {pe}"

        elif suffix == ".xlsx":
            try:
                import openpyxl
                wb = openpyxl.load_workbook(str(path), data_only=True)
                sheet_texts = []
                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    rows = []
                    for row in ws.iter_rows(values_only=True):
                        if any(v is not None for v in row):
                            rows.append(" | ".join(str(v) if v is not None else "" for v in row))
                    if rows:
                        sheet_texts.append(f"Sheet: {sheet_name}\n" + "\n".join(rows))
                return "\n\n".join(sheet_texts).strip() if sheet_texts else f"[Empty spreadsheet: {filename}]"
            except Exception as xe:
                return f"Error reading xlsx file: {xe}"

        else:
            # Unknown extensions: sniff for binary content (NUL bytes in the
            # header) instead of decoding it into token-exploding garbage.
            raw = path.read_bytes()
            if b"\x00" in raw[:4096]:
                return (
                    f"[Binary file: {Path(path).name}] This file is not text "
                    f"({len(raw)} bytes) and cannot be summarized as content."
                )
            return raw.decode("utf-8", errors="replace")

    except Exception as e:
        return f"Error: Could not read file: {e}"


def write_file(filename: str, content: str) -> str:
    """
    Write content to a file in the sandboxed directory.

    Args:
        filename: Relative path within the sandbox (no directory traversal allowed).
        content: String content to write.

    Returns:
        Success or error message.
    """
    try:
        path = _safe_resolve(filename, for_write=True)
    except ValueError as e:
        return f"Error: {e}"

    # Ensure parent directory exists
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
    except Exception as e:
        return f"Error: Could not create directory: {e}"

    try:
        path.write_text(content, encoding="utf-8")
        return f"Success: File written to {filename}"
    except Exception as e:
        return f"Error: Could not write file: {e}"
