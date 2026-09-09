"""
Test suite for all 6 root-cause fixes:
1. Citation number verification & grounding threshold (SOP-44 5mm vs 1/8 inch)
2. Multi-part query reasoning timeout & partial preservation
3. Sandbox injection hard refusal & audit logging
4. Deliverable read-back path resolution (outputs fallback, path normalization, pptx/xlsx support)
5. Sovereignty egress diagnostics (PID, process, cmdline, root-cause evidence)
6. Citation/list formatting collision prevention
"""

import os
import pytest
from pathlib import Path

from backend.tools.citation_tagger import tag_citations, strip_citations, _find_best_match
from backend.tools.calculator import solve_expression
from backend.core.audit_log import AuditLogger
from backend.tools.file_io import read_file, write_file, _safe_resolve, BASE_DIR, OUTPUT_DIR, WORKSPACE_DIR
from backend.agents.graph import _is_multipart_query, _split_multipart_query, synthesize_node
from backend.infra.sentinel_runner import SovereignSentinel


class TestFix1And6CitationGroundingAndFormatting:
    def test_sop44_number_grounding_verification(self):
        """Issue 1: 5mm is grounded in sop-44.txt, but 1/8 inch and 50mm are rejected."""
        sources = [
            {
                "text": "SOP-44: The maximum allowable corrosion depth for pressure vessels is 5mm.",
                "metadata": {"source": "sop-44.txt", "page": 1},
            }
        ]

        # Valid sentence with 5mm should be tagged
        s_valid = "The maximum allowable corrosion depth for pressure vessels is 5mm."
        tagged_valid = tag_citations(s_valid, sources)
        assert "[Source: sop-44.txt" in tagged_valid, f"Expected citation in: {tagged_valid}"

        # Hallucinated value (1/8 inch) MUST NOT be tagged with sop-44
        s_hallucinated = "The maximum allowable corrosion depth for pressure vessels is 1/8 inch."
        tagged_hallucinated = tag_citations(s_hallucinated, sources)
        assert "[Source:" not in tagged_hallucinated, f"Fabricated citation attached to 1/8 inch: {tagged_hallucinated}"

        # Hallucinated value (50mm) MUST NOT be tagged with sop-44
        s_wrong_number = "The maximum allowable corrosion depth for pressure vessels is 50mm."
        tagged_wrong_number = tag_citations(s_wrong_number, sources)
        assert "[Source:" not in tagged_wrong_number, f"Fabricated citation attached to 50mm: {tagged_wrong_number}"

    def test_list_formatting_collision_prevention(self):
        """Issue 6: Numbered list markers and citations must never collide."""
        sources = [
            {
                "text": "The maximum allowable corrosion depth for pressure vessels is 5mm.",
                "metadata": {"source": "sop-44.txt", "page": 1},
            }
        ]

        # Input with numbered list item
        text = "1. **Corrosion Limits:** The maximum allowable corrosion depth for pressure vessels is 5mm.\n2. **Frequency:** Inspection is quarterly."
        tagged = tag_citations(text, sources)

        # Line 1 citation must be after content, NOT between "1." and "**Corrosion Limits:**"
        lines = tagged.split("\n")
        assert len(lines) == 2, f"Line breaks were lost: {tagged}"
        assert lines[0].startswith("1. **Corrosion Limits:**"), f"Prefix was mangled: {lines[0]}"
        assert lines[0].endswith("]"), f"Citation must be appended at the end of the line: {lines[0]}"
        assert "1. [Source:" not in lines[0], f"Citation collided with list marker: {lines[0]}"

    def test_model_hallucinated_inline_citation_sanitization(self):
        """Issue 6: If model outputs 1. [Source: ...] **Header:**, it is properly normalized."""
        sources = [
            {
                "text": "The maximum allowable corrosion depth for pressure vessels is 5mm.",
                "metadata": {"source": "sop-44.txt", "page": 1},
            }
        ]
        text_with_bad_marker = "1. [Source: sop-44.txt, Page 1] **Corrosion Limits:** The maximum allowable corrosion depth for pressure vessels is 5mm."
        tagged = tag_citations(text_with_bad_marker, sources)
        assert tagged.startswith("1. **Corrosion Limits:**")
        assert "1. [Source:" not in tagged
        assert tagged.endswith("]")


class TestFix2MultiPartQueryDecomposition:
    def test_multipart_query_detection(self):
        """Issue 2: Correctly identifies multi-part queries."""
        q_single = "What is the corrosion limit in SOP-44?"
        assert not _is_multipart_query(q_single)

        q_numbered = """1. What is the corrosion limit?
2. What is the inspection frequency?
3. Who is the author?"""
        assert _is_multipart_query(q_numbered)
        parts = _split_multipart_query(q_numbered)
        assert len(parts) == 3
        assert parts[0]["header"] == "1."
        assert "corrosion limit" in parts[0]["prompt"]
        assert parts[1]["header"] == "2."
        assert "frequency" in parts[1]["prompt"]

    def test_multipart_subcall_partial_preservation(self):
        """Issue 2: Completed partial answers are preserved in synthesize_node."""
        state = {
            "input": """1. What is the corrosion limit?
2. What is the inspection frequency?""",
            "context": {},
            "retrieved_sources": [
                {"text": "Corrosion limit is 5mm. Inspection is quarterly.", "metadata": {"source": "sop-44.txt"}}
            ],
            "retrieval_invoked": True,
            "selected_model": "mock",
        }
        res = synthesize_node(state)
        assert "output" in res
        output = res["output"]
        assert "### 1." in output
        assert "### 2." in output


class TestFix3SandboxInjectionHandling:
    def test_calculator_blocked_injection_audit_and_refusal(self):
        """Issue 3: Calculator code execution attempts log audit event and surface hard refusal."""
        audit = AuditLogger()
        before_entries = len(audit.read_all_entries())

        # Attempt injection via calculator
        res = solve_expression("__import__(os).system(ls)")
        assert "[SECURITY_BLOCK]" in res

        # Verify audit log recorded SANDBOX_BLOCKED_INJECTION
        after_entries = audit.read_all_entries()
        assert len(after_entries) > before_entries
        last_entry = after_entries[-1]
        assert last_entry["event_type"] == "SANDBOX_BLOCKED_INJECTION"
        assert last_entry["details"]["tool"] == "calculator"
        assert "__import__" in last_entry["details"]["matched_pattern"]

    def test_synthesize_node_hard_refusal_on_security_block(self):
        """Issue 3: synthesize_node immediately refuses without generating workarounds."""
        state = {
            "input": "Calculate __import__(os).system(ls)",
            "context": {
                "step_0_result": "[SECURITY_BLOCK] Code execution or system access attempts via calculator are prohibited."
            },
            "retrieved_sources": [],
            "retrieval_invoked": False,
        }
        res = synthesize_node(state)
        assert "Execution blocked:" in res["output"]
        assert "os.popen" not in res["output"]
        assert res["model_used"] == "Security Policy (Sandbox)"


class TestFix4DeliverableReadBack:
    def test_read_deliverable_from_outputs_fallback(self):
        """Issue 4: file_io reads deliverables from workspace/outputs/ seamlessly."""
        # Write deliverable to outputs
        deliv_name = "test_audit_report.txt"
        deliv_path = OUTPUT_DIR / deliv_name
        deliv_path.write_text("Sovereign AI Deliverable Summary", encoding="utf-8")

        # 1. Read by bare filename
        content = read_file(deliv_name)
        assert content == "Sovereign AI Deliverable Summary"

        # 2. Read with outputs/ prefix
        assert read_file(f"outputs/{deliv_name}") == "Sovereign AI Deliverable Summary"

        # 3. Read with absolute path
        assert read_file(str(deliv_path.resolve())) == "Sovereign AI Deliverable Summary"

        deliv_path.unlink(missing_ok=True)

    def test_pptx_and_xlsx_read_support(self):
        """Issue 4: file_io can parse .pptx and .xlsx files."""
        import pptx
        import openpyxl

        # Test pptx
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Audit Architecture Slide"
        pptx_path = OUTPUT_DIR / "audit_slide.pptx"
        prs.save(str(pptx_path))

        pptx_text = read_file("audit_slide.pptx")
        assert "Audit Architecture Slide" in pptx_text

        # Test xlsx
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Metrics"
        ws.append(["Metric", "Value"])
        ws.append(["Uptime", "99.9%"])
        xlsx_path = OUTPUT_DIR / "audit_metrics.xlsx"
        wb.save(str(xlsx_path))

        xlsx_text = read_file("audit_metrics.xlsx")
        assert "Metric | Value" in xlsx_text
        assert "Uptime | 99.9%" in xlsx_text

        pptx_path.unlink(missing_ok=True)
        xlsx_path.unlink(missing_ok=True)


class TestFix5SovereigntyEgressDiagnostics:
    def test_synthetic_leak_diagnostics(self):
        """Issue 5: Synthetic probe captures PID, process, target, and root cause evidence."""
        sentinel = SovereignSentinel()
        diag = sentinel.trigger_synthetic_leak()

        assert "initiating_pid" in diag
        assert diag["initiating_pid"] == os.getpid()
        assert "initiating_process_name" in diag
        assert "protocol" in diag
        assert diag["protocol"] == "tcp"
        assert "destination_ip" in diag
        assert diag["destination_ip"] == "8.8.8.8"
        assert "destination_port" in diag
        assert diag["destination_port"] == 53
        assert "iptables_active" in diag
        assert "root_cause" in diag
        assert "active_background_telemetry_detected" in diag
        assert diag["active_background_telemetry_detected"] is False

        # Verify audit logger recorded SYNTHETIC_LEAK_TEST event with diagnostics
        audit = AuditLogger()
        entries = audit.read_all_entries()
        synth_entries = [e for e in entries if e.get("event_type") == "SYNTHETIC_LEAK_TEST"]
        assert len(synth_entries) > 0
        last_synth = synth_entries[-1]
        assert last_synth["details"]["initiating_pid"] == os.getpid()
        assert "root_cause" in last_synth["details"]